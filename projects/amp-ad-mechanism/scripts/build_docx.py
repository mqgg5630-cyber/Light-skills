#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成抗菌肽—AD 机制关联的文字稿 DOCX（纯文档，正文为主，无摘要与附录）。

运行：PYTHONUTF8=1 python3 build_docx.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"
OUT = ROOT / "outputs"
OUT.mkdir(parents=True, exist_ok=True)

SONG, HEI, TNR = "宋体", "黑体", "Times New Roman"


def set_run(run, *, cn=SONG, en=TNR, size=12, bold=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)


def p(doc, text="", *, cn=SONG, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
      indent=2.0, before=0, after=6, line=1.5):
    par = doc.add_paragraph()
    par.alignment = align
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(size * indent)
    if text:
        set_run(par.add_run(text), cn=cn, size=size, bold=bold)
    return par


def head(doc, text, size=13):
    return p(doc, text, cn=HEI, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
             indent=0, before=12, after=6, line=1.4)


def sub(doc, text):
    return p(doc, text, cn=HEI, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
             indent=0, before=8, after=4, line=1.4)


def ref(doc, idx, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Pt(24)
    pf.first_line_indent = Pt(-24)
    pf.space_after = Pt(3)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.3
    set_run(par.add_run(f"[{idx}] {text}"), size=10.5)
    return par


def build() -> Path:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = TNR
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.8)

    p(doc, "抗菌肽与阿尔茨海默病的机制关联：可用的机制落点与计算方案", cn=HEI, size=15, bold=True,
      align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, after=14, line=1.4)

    # 1
    head(doc, "一、目前的两条机制和它们的问题")
    p(doc, "现有的机制落点有两条。第一条是 AChE–Aβ 复合物的分子动力学：Aβ 对接进乙酰胆碱酯酶的外周阴离子位点（PAS）后，"
           "复合物在 1 μs 模拟中保持稳定，除 PAS 之外还有多处接触，Aβ 在 AChE 表面的主要停留区段是 344—361，"
           "这一段紧邻 PAS，但不受双位点抑制剂的空间位阻。它的实验基础是清楚的，AChE 本来就能通过 PAS 附近的疏水区加速 Aβ 成纤维，"
           "并形成毒性高于 Aβ 单独存在时的复合物[1-3]。第二条是抗菌肽与 Aβ 的交叉成核：两类肽都倾向形成 β-折叠，结构上兼容，"
           "可以互相成核，方向由界面和浓度决定，文献里既有加速的例子，也有抑制、改变纤维形态的例子[4-6]。")
    p(doc, "这两条本身没问题，问题在于它们撑不起一篇论文的机制部分。一是层级单一，两条都停在“肽与蛋白、肽与肽的一次相互作用”上，"
           "没有涉及受体与信号、膜界面、清除代谢、屏障转运和菌群，看上去就是一次对接加一段 MD。"
           "二是方向不定，交叉成核既能加速也能抑制，如果不事先规定判定方向的量化标准，得到什么结果都能解释一遍，等于不可证伪。"
           "三是没有交代暴露，抗菌肽是从肠道宏基因组里筛出来的，凭什么出现在 AChE 和 Aβ 所在的脑实质，这一步不回答，前面做得再细也会被质疑。")

    # 2
    head(doc, "二、别的疾病是怎么用对接和模拟把分子和疾病连起来的")
    p(doc, "把“某个分子和某个疾病有关”做成机制工作，别的方向已经有成熟套路，共同点不是“做了对接和 MD”，"
           "而是先指认疾病过程中的某个关键界面或决速步，再证明这个分子能在生理可及的浓度下改变它的方向或速率，"
           "最后给出体外、体内可以测量的对应量。常见的有五类。")
    p(doc, "一类是淀粉样蛋白之间的共聚集和交叉成核。Aβ 与 hIAPP 的工作是典型：用离散分子动力学采样异质寡聚体，"
           "统计异质接触是否多于同质接触、是否出现 β-桶中间体，再分别把单体放到纤维的生长端和侧表面，比较两种结合模式的自由能，"
           "结论是生长端结合优于侧表面结合，两种肽互为种子[7,8]。帕金森方向则用 α-syn 与细菌功能性淀粉样蛋白 CsgA 做同样的事[9-12]。")
    p(doc, "第二类是先天免疫受体识别。Aβ 原纤维与 CD14、TLR2、TLR4、RAGE 的识别就是用 MD 先取纤维低能构象、再做蛋白–蛋白对接，"
           "比较不同纤维多态和不同 pH 下的界面松紧[13]。FPR2 的冷冻电镜结构更直接，Aβ42 与 humanin 在 FPR2 上的结合位点高度重叠，"
           "界面面积分别是 1166 Å² 和 1077 Å²，humanin 正是靠占住同一口袋起保护作用[14]。")
    p(doc, "第三类是膜界面。抗菌肽本来就是靠膜起作用的，把神经元膜（含 GM1、胆固醇、PS）和细菌膜放在一起比较结合自由能，"
           "就能把“选择性丧失”量化出来；AMP 在有脂和无脂环境下还会在 cross-α 与 cross-β 之间切换[15,16]。")
    p(doc, "第四类是酶与底物的竞争。Aβ 在体内主要靠 IDE 和 NEP 降解，胰岛素等竞争性底物会拖慢这一过程[17,18]，"
           "做法是把候选分子和 Aβ 分别对接进催化腔，比停留时间和出口阻塞，再用竞争模型换算成清除速率的变化。")
    p(doc, "第五类是屏障转运和系统暴露，序列层面先做血脑屏障穿透预测，再算跨膜自由能或找受体介导转胞吞的结合面，"
           "最后用房室模型估算目标部位的游离浓度[19,20]。")
    p(doc, "还有一点要说清楚：本课题不按“抗菌肽作为抗 AD 抑制剂”来写。同样的计算，换个问法就从药物筛选变成机制研究——"
           "不问“能不能降低聚集总量”，而问“会不会改变聚集路径的分支比例”；不把 AChE 的 PAS 当成待抑制的口袋，而当成 Aβ 成核的模板表面；"
           "任何结论都附一句这个分子在该位点能达到多少浓度。LL-37 是现成的例子：它抑制长直纤维，却把可溶寡聚体稳住了，"
           "净效应偏向毒性而不是保护[4,5]；在 AD 脑内它升高，通过激活 CLIC1 造成小胶质细胞过度活化和神经炎症，"
           "在小鼠和猴上做出了 Aβ 升高、缠结增多、脑萎缩和认知损害的表型[6]。")

    # 3
    head(doc, "三、可以补上的机制落点")
    p(doc, "按上面五类范式，原来的两条可以扩成八个彼此独立又能交叉验证的落点，见图 1。前两个是原有机制的深化，后面六个是新增的层。")

    if (FIG / "fig1_mechanism_map.png").exists():
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.space_before = Pt(6)
        pic.paragraph_format.space_after = Pt(2)
        pic.add_run().add_picture(str(FIG / "fig1_mechanism_map.png"), width=Cm(15.4))
        p(doc, "图 1  肠道来源抗菌肽与 AD 之间的机制落点（按肠腔、循环与屏障、脑实质三个隔室排列）",
          cn=HEI, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, after=10, line=1.3)

    sub(doc, "1. AChE 的 PAS 上的三元复合物")
    p(doc, "把现有的 AChE–Aβ 体系改成三元的。抗菌肽如果能在 344—361 这段疏水区和 PAS 邻域稳定占位，"
           "结果有两种可能：一种是替代 Aβ 占位，减少可用的模板表面；另一种是形成 AChE–抗菌肽–Aβ 的夹层，反而把 Aβ 拉得更近。"
           "做法是以 PAS 残基（Tyr72、Asp74、Tyr124、Trp286、Tyr341）和 344—361 为活性残基做数据驱动对接，"
           "再跑二元与三元体系各三条 1 μs 轨迹，与已有的 AChE–Aβ 体系并排当内参。看的量是界面接触占据率、"
           "停留区段与 344—361 的重合比例，以及 Aβ 占据 PAS 的概率有没有变化。")

    sub(doc, "2. 与 Aβ 的交叉成核，重点是把方向定下来")
    p(doc, "交叉成核这条要改的不是方法而是判据。先用增强采样看异质二聚体里异质接触和同质接触的比例、有没有 β-桶中间体；"
           "再把抗菌肽分别放到 Aβ 纤维的生长端和侧表面，算结合自由能和伸长自由能，侧表面还要补一条伞形采样的自由能曲线，"
           "因为二次成核就发生在侧表面。方向在开跑前就规定好：生长端结合占优、伸长自由能更负，判为加速和共纤维化；"
           "侧表面结合占优并封端，判为改变形态、滞留寡聚体；两者都弱就是没有明显交叉成核。"
           "湿实验这边用 ThT 曲线做全局动力学拟合，把总效应拆成一次成核、二次成核和伸长三个速率常数[21,22]，"
           "和模拟的三种判断一一对应。另外 β-折叠含量对力场敏感，CHARMM36m 和 a99SB-disp 两套都要跑，只采信两套一致的结论[23,24]。")

    sub(doc, "3. tau 这一侧（可选）")
    p(doc, "以 tau 的 PHF6（306VQIVYK311）和 PHF6*（275VQIINK280）为界面做对接和采样，看阳离子肽能否靠芳香堆积参与异质 β-折叠。"
           "这一条阴性结果同样有用：如果抗菌肽因为同为阳离子而和 tau 相斥，正好说明效应是选择性地走 Aβ 和受体这两条线。")

    sub(doc, "4. 先天免疫受体")
    p(doc, "这是最容易和 AD 表型对上的一层，也是 LL-37 已经走通的一层。FPR2 这边可以直接利用 Aβ42 和 humanin 的复合物结构，"
           "把候选肽放进同一口袋做对接和膜环境下的 MD，看位点重合率、界面面积能否和 Aβ42 的 1166 Å² 同量级、"
           "以及 TM6 外移这类激活构象指标有没有出现[14]；CLIC1 这边按 LL-37 的路子算结合面，并在膜体系里看 CLIC1 跨膜段的插入倾向有没有变化[6]；"
           "TLR4–MD-2、CD14 和 RAGE 这边，把候选肽当作第二配体，评估它对 Aβ 原纤维–受体识别是竞争还是协同，"
           "顺便比较 pH 7.4 和炎症微环境 pH 6.0 的差别[13]。对应的细胞实验是小胶质细胞的炎症因子释放、NF-κB 报告和钙成像。")

    sub(doc, "5. 神经元膜界面")
    p(doc, "建三套膜：细菌内膜（POPE/POPG 3:1）、含 GM1 和胆固醇的神经元膜、内皮膜。先用 Martini 3 跑十几微秒看吸附和聚集，"
           "关键构象回到全原子验证。要的量是细菌膜与神经元膜的结合自由能差，也就是选择性，再加插入深度、膜厚和曲率扰动、水柱或孔的寿命。"
           "这里有一条独立于交叉成核的通道：如果抗菌肽在含 GM1 的膜面上富集，它不需要直接结合 Aβ，"
           "只要把 Aβ 招到同一个界面上就能促进成核，GM1 本来就是 Aβ 的成核平台[25]。")

    sub(doc, "6. Aβ 清除通路里的竞争")
    p(doc, "把候选肽和 Aβ 分别对接进 IDE 的封闭催化腔，比停留时间和腔口闭合程度，再用竞争模型把占位时间比换算成 Aβ 降解速率的相对下降；"
           "NEP 同理[17,18]。外周这一侧也值得算：如果候选肽与转甲状腺素蛋白、白蛋白的结合面和 Aβ 的结合面重叠，"
           "外周对 Aβ 的“沉降槽”作用就被削弱，不进脑也能把脑内 Aβ 负荷抬上去[26,27]。")

    sub(doc, "7. 抗菌肽自己的淀粉样化")
    p(doc, "抗菌肽和淀粉样蛋白共用一套结构语法。uperin 3.5 在无脂环境里形成 cross-β 纤维，遇到细菌膜脂就切换成 cross-α[15,16]，"
           "LL-37 的 17–29 片段能自组装成有功能的超分子纤维[28]。所以候选肽本身可能先在肠腔或循环里形成种子，"
           "再以种子而不是单体的形式参与宿主蛋白的异质成核。流程是先用聚集倾向预测找热点，枚举可能的 steric zipper，"
           "对最优的做几百纳秒稳定性模拟（看剥离功和氢键网络），再把稳定的 zipper 当作种子面去算 Aβ 的结合和伸长自由能。"
           "如果这一条成立，机制就变成单体通道加种子通道两条，叙述上也更完整。")

    sub(doc, "8. 肠道这一端")
    p(doc, "肠道端本身就能构成完整链条，不必先进脑。curli 的主亚基 CsgA 是典型的 cross-β 结构，它的片段种子能促进 Aβ 成纤维[9]；"
           "动物实验里产 curli 的细菌会增强肠和脑内 α-syn 的聚集，并放大 TLR2、IL-6、TNF 的反应[10,11]；"
           "来自人体菌群的多种 CsgA 同源物也能以 1:1 复合物的形式加速 α-syn 聚集[12]。抗菌肽进到这个系统里有两个可以算的面。"
           "分子面是它与 CsgA、FapC 组装界面的相互作用，是封端、还是打乱重复单元的配准、还是把可溶 CsgA 稳成寡聚体——"
           "注意方向，抑制 curli 成纤维未必是好事，把胞外纤维变成可溶寡聚体，反而可能提高跨上皮和交叉成核的能力。"
           "群落面是用候选肽的 MIC 谱作约束，做菌群代谢的群落模型，预测短链脂肪酸产生菌和高产 LPS 菌的相对变化，"
           "再顺下去看屏障完整性和外周炎症[29,30]。这条链的读数是：MIC 谱、群落模型给出的丁酸丙酸通量和 LPS 负荷、"
           "肠屏障跨上皮电阻、血清 LPS 与细胞因子、小胶质表型。")

    # 4
    head(doc, "四、肠道来源意味着什么：要不要过血脑屏障")
    p(doc, "结论先放：对肠道宏基因组来源的抗菌肽来说，“必须穿过血脑屏障才能影响大脑”不是默认前提，而是三条路径里最苛刻的一条。"
           "建议把肠道局部作用和外周作用当主线，把跨屏障入脑当成需要额外条件的分支。")
    p(doc, "三条路径是这样分的。第一条是直接入脑，肽经肠上皮吸收入血，躲过蛋白酶和肾清除，再靠被动扩散或受体介导转胞吞过血脑屏障，"
           "在脑间质达到能和 AChE、Aβ、受体作用的游离浓度。第二条是在肠道就地起效，肠腔里的浓度可以到微摩尔甚至更高，"
           "作用对象是菌群、细菌功能性淀粉样蛋白、肠上皮和肠神经系统，再经迷走传入和免疫细胞、细胞因子把信号送到中枢。"
           "第三条是入血但不进脑，改变外周对 Aβ 的结合与清除（转甲状腺素蛋白、白蛋白、单核巨噬），或者影响 Aβ 跨屏障的净通量。")
    p(doc, "之所以不把第一条当主线，是因为细菌源候选肽大多十几到几十个残基、净正电荷高、两亲性强，这类分子被动过屏障的效率很低；"
           "而能过血脑屏障的肽在序列特征上自成一类——偏小、偏疏水、弱阳离子，和一般的穿膜肽并不是一回事[19,20]。"
           "与此同时，肠道端和外周端已经有被验证过的完整链条可以用[10-12,26,27,29,30]。")
    p(doc, "如果还是要主张能入脑，需要同时满足四条：序列层面的血脑屏障穿透预测为阳性，且不只依赖一个模型，并报出概率而不是只给标签；"
           "跨膜自由能垒处在可跨越的范围，或者能明确指出受体介导转胞吞的结合面；房室模型给出的脑间质游离浓度不低于目标结合常数的十分之一；"
           "脑组织里检出的是完整肽，用质谱定量，而不是降解片段或者只看荧光标记。四条里有一条不成立，就把机制主线放回肠道端和外周端。"
           "这不是退让，脑内 Aβ 稳态本来就受外周清除和系统炎症的调控[26,27]，肠源因子不进脑也能改变中枢的病理负荷。")
    p(doc, "还有一个低成本但关键的实验：候选肽在模拟肠液和血浆里的稳定性以及降解片段谱。如果主要活性片段变短、变疏水，"
           "入脑的可能性反而上升，这时脑内那几个落点就该拿片段而不是全长肽去建模。")

    # 5
    head(doc, "五、计算方案与判据")
    p(doc, "八个落点共用一套流程，评审看的不是某一次对接分数，而是整条链条是否自洽、能否复现，所以结构来源、力场、时长、判据和阴性对照要统一。")
    p(doc, "序列层面先做分诊：净电荷、疏水矩、两亲性，聚集倾向预测，血脑屏障穿透预测，以及溶血和细胞毒预测，"
           "同时给出各隔室的可及浓度估计——这一步不过关，后面的结论都只在试管里成立。结构用 ESMFold 或 AlphaFold3 生成，"
           "无序肽保留构象系综而不是单一构象，并准备 pH 7.4 和 pH 6.0 两套质子化状态。")
    p(doc, "对接只用来枚举可能的界面，不进入结论句。数据驱动对接（HADDOCK）和盲对接（HPEPDOCK、ClusPro）并用，"
           "再用 AlphaFold-Multimer 交叉验证，要求至少两个引擎给出一致的位姿、首位聚类占比不低于三成。")
    p(doc, "动力学部分，全原子体系每个跑三条 1 μs 的独立轨迹，310 K、0.15 mol/L NaCl；膜体系先用 Martini 3 跑长时程，"
           "关键构象回到全原子；自由能用伞形采样或 metadynamics，MM/PBSA 只做相对排序，不和实验亲和力直接比。"
           "判定稳定的标准是最后 500 ns 的界面接触占据率不低于六成，而且至少两条副本一致。")
    p(doc, "阴性对照必须和正样本同流程、同时长：打乱序列的肽、等长等电荷的非抗菌肽、同源但没有抗菌活性的肽，这三类都应当通不过上面的判据。"
           "能做到这一点，“为什么是这条肽而不是随便哪条阳离子肽”这个问题才算回答了。")

    # 6
    head(doc, "六、接下来的安排")
    p(doc, "优先做两件事。一是把肠道端（第 8 条）和外周清除竞争（第 6 条）立起来，这两条不依赖跨屏障假设，风险最低，"
           "而且能直接给出“肠道来源”这个出身的意义。二是把原有的 AChE 三元体系和交叉成核的方向判据补完，"
           "这两条是已经动手的部分，补上三元体系、双力场和 ThT 全局拟合之后，机制层次和可证伪性都能上一个台阶。")
    p(doc, "受体那一层（第 4 条）放在第二批，它最容易和 AD 表型对上，但需要膜体系和细胞实验配合，周期较长。"
           "膜界面（第 5 条）、自淀粉样化（第 7 条）、tau（第 3 条）作为第三批，视前两批的结果决定是否展开。"
           "跨血脑屏障那条路径在稳定性和降解片段谱的数据出来之前，先按条件分支处理，不写进主线结论。")

    # 参考文献
    head(doc, "参考文献")
    refs = [
        "Inestrosa N C, Alvarez A, Pérez C A, et al. Acetylcholinesterase accelerates assembly of amyloid-β-peptides into Alzheimer’s fibrils[J]. Neuron, 1996, 16(4): 881–891.",
        "De Ferrari G V, Canales M A, Shin I, et al. A structural motif of acetylcholinesterase that promotes amyloid β-peptide fibril formation[J]. Biochemistry, 2001, 40(35): 10447–10457.",
        "AChE–Aβ 复合物的 1 μs 分子动力学研究（原稿编号 [31] 的文献，著录信息待补）。",
        "De Lorenzi E, Chiari M, Colombo R, et al. Evidence that the human innate immune peptide LL-37 may be a binding partner of amyloid-β and inhibitor of fibril assembly[J]. Journal of Alzheimer’s Disease, 2017, 59(4): 1213–1226.",
        "LL-37 and its truncated fragments modulate amyloid-β dynamics, aggregation and toxicity through hetero-oligomer and cluster formation[J]. 2025. PMID: 40916348.",
        "Wang C, et al. Human antimicrobial peptide LL-37 contributes to Alzheimer’s disease progression[J]. Molecular Psychiatry, 2022. doi:10.1038/s41380-022-01790-6.",
        "Fan X, Zhang X, Yan J, et al. Computational investigation of co-aggregation and cross-seeding between Aβ and hIAPP underpinning the crosstalk in Alzheimer’s disease and type-2 diabetes[J]. Journal of Chemical Information and Modeling, 2024, 64(13): 5303–5316.",
        "Identification of hybrid amyloid strains assembled from amyloid-β and human islet amyloid polypeptide[J]. Nanotechnology, 2023. doi:10.1088/1361-6528/acf3ee.",
        "Perov S, Lidor O, Salinas N, et al. Structural insights into curli CsgA cross-β fibril architecture inspire repurposing of anti-amyloid compounds as anti-biofilm agents[J]. PLoS Pathogens, 2019, 15(8): e1007978.",
        "Chen S G, Stribinskis V, Rane M J, et al. Exposure to the functional bacterial amyloid protein curli enhances alpha-synuclein aggregation in aged Fischer 344 rats and Caenorhabditis elegans[J]. Scientific Reports, 2016, 6: 34477.",
        "Wang C, Lau C Y, Ma F, et al. Genome-wide screen identifies curli amyloid fibril as a bacterial component promoting host neurodegeneration[J]. PNAS, 2021, 118(34): e2106504118.",
        "Bhoite S S, Han Y, Ruotolo B T, et al. Mechanistic insights into accelerated α-synuclein aggregation mediated by human microbiome-associated functional amyloids[J]. Journal of Biological Chemistry, 2022, 298(8): 102088.",
        "Factors driving amyloid beta fibril recognition by cell surface receptors: a computational study[J]. 2025. PMCID: PMC12566521.",
        "Zhu Y, Lin X, Zong X, et al. Structural basis of FPR2 in recognition of Aβ42 and neuroprotection by humanin[J]. Nature Communications, 2022, 13: 1775.",
        "Salinas N, Tayeb-Fligelman E, Sammito M D, et al. The amphibian antimicrobial peptide uperin 3.5 is a cross-α/cross-β chameleon functional amyloid[J]. PNAS, 2021, 118(3): e2014442118.",
        "Bücker R, Seuring C, Cazey C, et al. The cryo-EM structures of two amphibian antimicrobial cross-β amyloid fibrils[J]. Nature Communications, 2022, 13: 4356.",
        "Farris W, Mansourian S, Chang Y, et al. Insulin-degrading enzyme regulates the levels of insulin, amyloid β-protein, and the β-amyloid precursor protein intracellular domain in vivo[J]. PNAS, 2003, 100(7): 4162–4167.",
        "Iwata N, Tsubuki S, Takaki Y, et al. Metabolic regulation of brain Aβ by neprilysin[J]. Science, 2001, 292(5521): 1550–1552.",
        "Kumar V, Patiyal S, Dhall A, et al. B3Pred: a random-forest-based method for predicting and designing blood–brain barrier penetrating peptides[J]. Pharmaceutics, 2021, 13(8): 1237.",
        "Prediction of blood-brain barrier-penetrating peptides using B3BPFN[J]. Frontiers in Molecular Biosciences, 2026. doi:10.3389/fmolb.2026.1858506.",
        "Meisl G, Kirkegaard J B, Arosio P, et al. Molecular mechanisms of protein aggregation from global fitting of kinetic models[J]. Nature Protocols, 2016, 11(2): 252–272.",
        "Cohen S I A, Linse S, Luheshi L M, et al. Proliferation of amyloid-β42 aggregates occurs through a secondary nucleation mechanism[J]. PNAS, 2013, 110(24): 9758–9763.",
        "Huang J, Rauscher S, Nawrocki G, et al. CHARMM36m: an improved force field for folded and intrinsically disordered proteins[J]. Nature Methods, 2017, 14(1): 71–73.",
        "Robustelli P, Piana S, Shaw D E. Developing a molecular dynamics force field for both folded and disordered protein states[J]. PNAS, 2018, 115(21): E4758–E4766.",
        "Matsuzaki K. Aβ–ganglioside interactions in the pathogenesis of Alzheimer’s disease[J]. Biochimica et Biophysica Acta – Biomembranes, 2020, 1862(8): 183233.",
        "Wang J, Gu B J, Masters C L, et al. A systemic view of Alzheimer disease—insights from amyloid-β metabolism beyond the brain[J]. Nature Reviews Neurology, 2017, 13(10): 612–623.",
        "Xiang Y, Bu X L, Liu Y H, et al. Physiological amyloid-beta clearance in the periphery and its therapeutic potential for Alzheimer’s disease[J]. Acta Neuropathologica, 2015, 130(4): 487–499.",
        "Engelberg Y, Landau M. The human LL-37(17-29) antimicrobial peptide reveals a functional supramolecular structure[J]. Nature Communications, 2020, 11: 3894.",
        "Diener C, Gibbons S M, Resendis-Antonio O. MICOM: metagenome-scale modeling to infer metabolic interactions in the gut microbiota[J]. mSystems, 2020, 5(1): e00606-19.",
        "Yang J, Liang J, Hu N, et al. The gut microbiota modulates neuroinflammation in Alzheimer’s disease: elucidating crucial factors and mechanistic underpinnings[J]. CNS Neuroscience & Therapeutics, 2024, 30(10): e70091.",
    ]
    for i, r in enumerate(refs, 1):
        ref(doc, i, r)

    out = OUT / "抗菌肽-阿尔茨海默病-机制关联框架.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    print("wrote", build())
